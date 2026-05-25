"""
Base Excel generator for filling templates.
Loads a template and fills cells according to a JSON mapping.
Preserves images and formatting automatically with openpyxl + Pillow.
"""
import json
from pathlib import Path
from typing import Any
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.workbook import Workbook


# Base directories
BASE_DIR = Path(__file__).resolve().parent.parent.parent  # → app/
TEMPLATES_DIR = BASE_DIR / "templates_xlsx"
MAPPING_DIR = Path(__file__).resolve().parent / "mapping"


class ExcelGeneratorBase:
    """
    Base Excel generator.
    Subclasses only need to define `mapping_file`.
    """

    mapping_file: str = ""

    def __init__(self):
        self.mapping = self._load_mapping()
        self.template_path = TEMPLATES_DIR / self.mapping["template_file"]

        if not self.template_path.exists():
            raise FileNotFoundError(
                f"Template not found: {self.template_path}"
            )

    def _load_mapping(self) -> dict:
        """Load the JSON mapping file."""
        mapping_path = MAPPING_DIR / self.mapping_file
        if not mapping_path.exists():
            raise FileNotFoundError(f"Mapping not found: {mapping_path}")

        with open(mapping_path, "r", encoding="utf-8") as f:
            return json.load(f)

    def generate(self, form_data: dict) -> BytesIO:
        """
        Generate the Excel file with form data.

        Args:
            form_data: dictionary with form data

        Returns:
            BytesIO with Excel file content
        """
        wb = load_workbook(self.template_path)
        self._fill_sheets(wb, form_data)
        self._insert_photos(wb, form_data)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def _fill_sheets(self, wb: Workbook, form_data: dict):
        """Iterate through all mapped sheets and fill cells."""
        for sheet_name, sections in self.mapping["sheets"].items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]

            for section_name, fields in sections.items():
                section_data = form_data.get(section_name, {})
                for field_name, cell_ref in fields.items():
                    value = section_data.get(field_name)
                    if value is not None and value != "":
                        self._write_cell(ws, cell_ref, value)

    def _write_cell(self, ws, cell_ref: str, value: Any):
        """
        Write to a cell, handling merged cells.
        For merged cells, write to the top-left cell of the range.
        """
        try:
            ws[cell_ref] = value
        except AttributeError:
            # Cell is part of a merged range but not the top-left corner
            print(f"⚠️  Cannot write to {cell_ref} (merged cell): {value}")
            # Find the root cell of the merge
            for merged_range in ws.merged_cells.ranges:
                if cell_ref in merged_range:
                    top_left = merged_range.start_cell
                    print(f"   Using instead: {top_left.coordinate}")
                    top_left.value = value
                    break

    def _insert_photos(self, wb: Workbook, form_data: dict):
        """
        Insert photos into photographic record sheets.
        Reads photo sheet names from the JSON mapping for flexibility.
        """
        from openpyxl.drawing.image import Image as XLImage
        from pathlib import Path
        
        # Get photo sheets from mapping (defaults to RI sheet names if not defined)
        photo_sheets = self.mapping.get("photo_sheets", [
            ["REGISTRO FOTOGRÁFICO", "fotos_general"],
            ["REGISTRO FOTOGRÁFICO HALLAZGOS", "fotos_hallazgos"],
        ])
        
        # 12 photo positions: (photo_cell, description_cell, width, height)
        positions = [
            ("A7", "A25", 200, 150),     # Photo 1
            ("C7", "C25", 200, 150),     # Photo 2
            ("A30", "A48", 200, 150),    # Photo 3
            ("C30", "C48", 200, 150),    # Photo 4
            ("A53", "A71", 200, 150),    # Photo 5
            ("C53", "C71", 200, 150),    # Photo 6
            ("A76", "A94", 200, 150),    # Photo 7
            ("C76", "C94", 200, 150),    # Photo 8
            ("A99", "A117", 200, 150),   # Photo 9
            ("C99", "C117", 200, 150),   # Photo 10
            ("A122", "A140", 200, 150),  # Photo 11
            ("C122", "C140", 200, 150),  # Photo 12
        ]
        
        for sheet_name, data_key in photo_sheets:
            if sheet_name not in wb.sheetnames:
                continue
            
            data = form_data.get(data_key)
            if not data:
                continue
            
            ws = wb[sheet_name]
            
            # Fill A5 with concatenated info
            site = data.get("sede_educativa", "")
            date = data.get("fecha", "")
            period = data.get("periodo", "")
            info_text = f"Sede educativa: {site}     Fecha: {date}     Período: {period}"
            ws["A5"] = info_text
            
            # Insert all 12 photos and descriptions
            for idx, (photo_cell, desc_cell, width, height) in enumerate(positions, 1):
                photo_path = data.get(f"foto_{idx}")
                if photo_path and Path(photo_path).exists():
                    img = XLImage(photo_path)
                    img.width = width
                    img.height = height
                    ws.add_image(img, photo_cell)
                
                description = data.get(f"descripcion_{idx}", "")
                if description:
                    ws[desc_cell] = f"Descripción F{idx}: {description}"